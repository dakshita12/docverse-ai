import fitz
from docx import Document
from pptx import Presentation
from io import BytesIO
from PIL import Image
from src.processing.text_processing import clean_text
from src.processing.ocr import OCRProcessor


MIN_TEXT_LENGTH = 50


def should_use_ocr(text: str) -> bool:
    """
    Determines whether OCR should be used for a PDF page.

    Args:
        text: Text extracted from the PDF page.

    Returns:
        True if OCR is required, otherwise False.
    """
    return len(text.strip()) < MIN_TEXT_LENGTH


def read_pdf(file_path):
    """
    Reads a PDF file and returns its text.
    """

    ocr = OCRProcessor()

    pages_text = []

    with fitz.open(file_path) as document:

        for page in document:
            page_text = page.get_text().strip()

            if should_use_ocr(page_text):
                pixmap = page.get_pixmap()
                img_bytes = pixmap.tobytes("png")
                image_stream = BytesIO(img_bytes)
                image = Image.open(image_stream)
                page_text = ocr.extract_text_from_image(image)
            
            pages_text.append(page_text)

    text = "\n".join(pages_text)

    text = clean_text(text)
    return text


def read_docx(file_path):
    """
    Reads a DOCX file and returns its text.
    """
    document = Document(file_path)

    text = "" 

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    text = clean_text(text)
    return text


def read_pptx(file_path):
    """
    Reads a PPTX file and returns its text.
    """
    presentation = Presentation(file_path)

    text = ""

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

    text = clean_text(text)
    return text


def load_document(file_path):
    """
    Loads a document based on its file extension.
    """

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        extension = file_path.suffix.lower()

        if extension == ".pdf":
            return read_pdf(file_path)

        elif extension == ".docx":
            return read_docx(file_path)

        elif extension == ".pptx":
            return read_pptx(file_path)

        raise ValueError(f"Unsupported file type: {extension}")

    except (FileNotFoundError, ValueError):
        raise 
    
    except Exception as e:
        raise RuntimeError(f"Error loading document: {e}")